"""HTTP API and the public unsubscribe endpoint."""

from __future__ import annotations

import datetime as dt
import threading
from dataclasses import dataclass

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import delete, func, select, update
from sqlalchemy.orm import Session

from app.config import settings
from app.db import get_db, session_scope
from app.models import (
    AIAnalysis, Campaign, CampaignRun, Company, Contact, EmailEvent,
    EmailMessage, Lead, Mailbox, SearchQuery, Signal, Source, Suppression,
    Service, WebsiteAudit, WhatsAppMessage, as_aware, utcnow,
)
from app.services import (
    ai, auth, billing, catalog, inbox, linkedin, mailer, notify, pipeline, scoring, tenant, whatsapp,
)
from app.services.compliance import COUNTRY_NAMES, COUNTRY_POLICY, policy_for_country, suppress
from app.services.discovery import available_sources
from app.services.signals import company_signals, signals_by_company

router = APIRouter(prefix="/api")

# Guards against two runs of the same campaign overlapping.
_running: set[int] = set()
_running_lock = threading.Lock()


# ---------------------------------------------------------------------------
# Who is asking. Every query below is limited to this user's own rows.
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Me:
    id: int
    role: str

    @property
    def is_admin(self) -> bool:
        return self.role == "admin"


def current_user(request: Request, db: Session = Depends(get_db)) -> Me:
    ident = getattr(request.state, "user", None)
    if ident:
        if ident.get("status") != "active":
            raise HTTPException(403, "Your account is suspended. Please contact support.")
        return Me(ident["id"], ident["role"])
    if not auth.login_required():
        # Local use with no ADMIN_PASSWORD: the one person at the keyboard is the admin.
        return Me(tenant.admin_id(db), "admin")
    raise HTTPException(401, "Not authenticated")


def admin_only(me: Me = Depends(current_user)) -> Me:
    if not me.is_admin:
        raise HTTPException(403, "Only the platform admin can do this")
    return me


def _owned(db: Session, model, row_id: int, me: Me, what: str = "item"):
    """Fetch a row that must belong to this user. Someone else's row is a 404,
    not a 403, so ids do not reveal what exists."""
    row = db.get(model, row_id)
    if row is None or row.owner_id != me.id:
        raise HTTPException(404, f"{what} not found")
    return row


def _require_bot(db: Session, me: Me, bot_key: str) -> None:
    if me.is_admin:
        raise HTTPException(403, "The admin account runs the platform and has no agents of its own. "
                                 "Sign in with a customer account to run bots.")
    if not billing.has_access(db, me.id, bot_key):
        raise HTTPException(402, f"You don't have access to the {catalog.label_for(bot_key)} bot. "
                                 f"Start a free trial or subscribe under My bots.")


# ---------------------------------------------------------------------------
# Meta & stats
# ---------------------------------------------------------------------------


@router.get("/meta")
def meta(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import User

    cfg = tenant.config_for(db, me.id)
    user = db.get(User, me.id)
    trial_end = None
    if user and user.trial_started_at:
        trial_end = (as_aware(user.trial_started_at) + dt.timedelta(hours=settings.trial_hours)).isoformat()
    bots = billing.catalog(db)
    return {
        "app": settings.platform_name,
        "user": {
            "id": me.id, "role": me.role, "name": user.name if user else "",
            "email": user.email if user else "", "company": user.company if user else "",
        },
        "access": billing.access_map(db, me.id),
        "trial": {"used": bool(user and user.trial_started_at),
                  "bot": user.trial_bot if user else None, "ends_at": trial_end,
                  "hours": settings.trial_hours},
        "bots": [billing.bot_json(b) for b in bots],
        "razorpay_ready": billing.razorpay_ready(),
        "ai_ready": ai.available(),
        "send_on_server": settings.send_on_server,
        "email_ready": cfg.smtp_ready,
        "sender_configured": cfg.sender_ready,
        "sources": available_sources(),
        "services": [{"key": x["key"], "label": x["label"], "builtin": x["builtin"]}
                     for x in catalog.all_services()],
        "country_policy": {
            code: {"name": COUNTRY_NAMES.get(code, code), "policy": p, "reason": r}
            for code, (p, r) in COUNTRY_POLICY.items()
        },
        "scheduler_enabled": settings.scheduler_enabled,
        "timezone": cfg.timezone,
        "telegram_ready": notify.ready() and me.is_admin,
        "inbox_ready": cfg.imap_ready,
        "whatsapp_enabled": settings.whatsapp_enabled,
        "support_email": settings.support_email,
    }


@router.get("/stats/summary")
def stats_summary(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    today = utcnow().replace(hour=0, minute=0, second=0, microsecond=0)

    def count(model, *where):
        q = select(func.count(model.id)).where(_mine(model, me))
        for w in where:
            q = q.where(w)
        return q.scalar_subquery()

    events = lambda kind: (  # noqa: E731 - events belong to the user through their message
        select(func.count(EmailEvent.id)).join(EmailMessage, EmailMessage.id == EmailEvent.message_id)
        .where(EmailMessage.owner_id == me.id, EmailEvent.type == kind).scalar_subquery())

    # One SELECT of subqueries: a single round trip instead of seventeen.
    parts = {
        "total_leads": count(Lead),
        "new_today": count(Lead, Lead.created_at >= today),
        "qualified": count(Lead, Lead.score >= scoring.MEDIUM_THRESHOLD, Lead.stage != "rejected"),
        "high_intent": count(Lead, Lead.intent_level == "high", Lead.stage != "rejected"),
        "with_email": count(Lead, Lead.contact_id.isnot(None)),
        "pending_approval": count(EmailMessage, EmailMessage.status == "pending_approval"),
        "queued": count(EmailMessage, EmailMessage.status == "queued"),
        "emails_sent": count(EmailMessage, EmailMessage.status == "sent"),
        "sent_today": count(EmailMessage, EmailMessage.status == "sent", EmailMessage.sent_at >= today),
        "replies": events("reply"),
        "bounces": events("bounce"),
        "unsubscribed": count(Suppression),
        "rejected": count(Lead, Lead.stage == "rejected"),
        "whatsapp_pending": count(WhatsAppMessage, WhatsAppMessage.status == "pending"),
        "interested": count(Lead, Lead.status == "interested"),
        "companies": select(func.count(func.distinct(Lead.company_id)))
                     .where(Lead.owner_id == me.id).scalar_subquery(),
        "ai_cost_usd": select(func.coalesce(func.sum(AIAnalysis.cost_usd), 0.0))
                       .join(Lead, Lead.id == AIAnalysis.lead_id)
                       .where(Lead.owner_id == me.id).scalar_subquery(),
    }
    row = db.execute(select(*[v.label(k) for k, v in parts.items()])).one()._mapping
    out = {k: row[k] for k in parts}
    out["ai_cost_usd"] = round(float(out["ai_cost_usd"] or 0), 4)
    out["daily_cap"] = tenant.config_for(db, me.id).daily_cap
    return out


def _mine(model, me: Me):
    return model.owner_id == me.id


@router.get("/stats/breakdown")
def stats_breakdown(by: str = "service", me: Me = Depends(current_user),
                    db: Session = Depends(get_db)):
    column = {
        "service": Lead.service_category,
        "source": Lead.source_kind,
        "intent": Lead.intent_level,
        "status": Lead.status,
        "stage": Lead.stage,
    }.get(by)
    if column is not None:
        rows = db.execute(
            select(column, func.count(Lead.id)).where(Lead.owner_id == me.id)
            .group_by(column).order_by(func.count(Lead.id).desc())
        ).all()
    elif by == "country":
        rows = db.execute(
            select(Company.country, func.count(Lead.id))
            .join(Lead, Lead.company_id == Company.id).where(Lead.owner_id == me.id)
            .group_by(Company.country).order_by(func.count(Lead.id).desc())
        ).all()
    elif by == "industry":
        rows = db.execute(
            select(Company.industry, func.count(Lead.id))
            .join(Lead, Lead.company_id == Company.id).where(Lead.owner_id == me.id)
            .group_by(Company.industry).order_by(func.count(Lead.id).desc()).limit(12)
        ).all()
    else:
        raise HTTPException(400, f"unknown breakdown '{by}'")
    return [{"key": k or "unknown", "count": c} for k, c in rows]


@router.get("/stats/timeseries")
def stats_timeseries(days: int = 30, me: Me = Depends(current_user),
                     db: Session = Depends(get_db)):
    since = utcnow() - dt.timedelta(days=days)
    leads = db.execute(
        select(func.date(Lead.created_at), func.count(Lead.id))
        .where(Lead.owner_id == me.id, Lead.created_at >= since).group_by(func.date(Lead.created_at))
    ).all()
    sent = db.execute(
        select(func.date(EmailMessage.sent_at), func.count(EmailMessage.id))
        .where(EmailMessage.owner_id == me.id, EmailMessage.sent_at >= since)
        .group_by(func.date(EmailMessage.sent_at))
    ).all()
    lead_map = {str(d): c for d, c in leads if d}
    sent_map = {str(d): c for d, c in sent if d}
    out = []
    for i in range(days, -1, -1):
        day = (utcnow() - dt.timedelta(days=i)).date().isoformat()
        out.append({"day": day, "leads": lead_map.get(day, 0), "sent": sent_map.get(day, 0)})
    return out


# ---------------------------------------------------------------------------
# Campaigns
# ---------------------------------------------------------------------------


class CampaignIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    service_category: str
    sources: list[str] = []
    keywords: list[str] = []
    locations: list[str] = []
    countries: list[str] = ["GB", "US"]
    schedule_time: str = "07:00"
    schedule_days: str = "mon-fri"
    enabled: bool = True
    daily_lead_limit: int = Field(25, ge=1, le=500)
    daily_email_limit: int = Field(20, ge=0, le=1000)
    approval_mode: str = "manual"
    min_score: int = Field(55, ge=0, le=100)
    rotate_cities: bool = True

    # A bad value here would not fail on save - it would fail at 07:00 when the
    # scheduler tries to build the trigger, silently skipping the campaign.
    @field_validator("name")
    @classmethod
    def _name(cls, v: str) -> str:
        if not v.strip():
            raise ValueError("name is required")
        return v.strip()

    @field_validator("service_category")
    @classmethod
    def _service(cls, v: str) -> str:
        if not catalog.exists(v):
            raise ValueError(f"unknown service '{v}'")
        return v

    @field_validator("schedule_time")
    @classmethod
    def _time(cls, v: str) -> str:
        try:
            hour, minute = (int(p) for p in v.split(":")[:2])
        except ValueError:
            raise ValueError("time must be HH:MM") from None
        if not (0 <= hour < 24 and 0 <= minute < 60):
            raise ValueError("time must be HH:MM")
        return f"{hour:02d}:{minute:02d}"

    @field_validator("schedule_days")
    @classmethod
    def _days(cls, v: str) -> str:
        from app.scheduler import DAY_MAP
        if v not in DAY_MAP:
            raise ValueError(f"days must be one of {', '.join(DAY_MAP)}")
        return v

    @field_validator("approval_mode")
    @classmethod
    def _approval(cls, v: str) -> str:
        if v not in ("manual", "auto"):
            raise ValueError("approval must be manual or auto")
        return v

    @field_validator("countries")
    @classmethod
    def _countries(cls, v: list[str]) -> list[str]:
        return list(dict.fromkeys(c.strip().upper() for c in v if c.strip()))

    @field_validator("sources", "keywords", "locations")
    @classmethod
    def _clean_list(cls, v: list[str]) -> list[str]:
        return list(dict.fromkeys(s.strip() for s in v if s.strip()))


class ServiceIn(BaseModel):
    label: str = Field(min_length=2, max_length=80)
    pitch: str = Field(min_length=10, max_length=600)
    keywords: list[str] = []
    intent_phrases: list[str] = []
    industries: list[str] = []

    @field_validator("label", "pitch")
    @classmethod
    def _text(cls, v: str) -> str:
        v = " ".join(v.split())
        if len(v) < 2:
            raise ValueError("this field is required")
        return v

    @field_validator("keywords", "intent_phrases", "industries")
    @classmethod
    def _terms(cls, v: list[str]) -> list[str]:
        seen: dict[str, str] = {}
        for item in v:
            item = " ".join(item.split())
            if item:
                seen.setdefault(item.lower(), item)
        if len(seen) > 30:
            raise ValueError("at most 30 entries")
        return list(seen.values())

    @field_validator("keywords")
    @classmethod
    def _need_keywords(cls, v: list[str]) -> list[str]:
        if not v:
            raise ValueError("add at least one keyword - it is what the agent searches for")
        return v


def _service_json(db: Session, svc: dict, me: Me) -> dict:
    used = db.execute(
        select(func.count(Campaign.id)).where(Campaign.service_category == svc["key"],
                                              Campaign.owner_id == me.id)
    ).scalar_one()
    return {**svc, "campaigns": used}


@router.get("/services")
def list_services(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.services.discovery.base import KEYWORD_GROUPS
    out = []
    for svc in catalog.all_services():
        if svc["builtin"]:
            svc = {**svc, "keywords": KEYWORD_GROUPS.get(svc["key"], []),
                   "pitch": ai.SERVICE_PITCH.get(svc["key"], ""),
                   "intent_phrases": [], "industries": []}
        out.append(_service_json(db, svc, me))
    return out


@router.post("/services")
def create_service(payload: ServiceIn, me: Me = Depends(admin_only), db: Session = Depends(get_db)):
    taken = set(db.execute(select(Service.key)).scalars())
    key = catalog.slugify(payload.label, taken)
    if any(x["label"].lower() == payload.label.lower() for x in catalog.all_services()):
        raise HTTPException(409, f"a service called '{payload.label}' already exists")
    db.add(Service(key=key, **payload.model_dump()))
    # Every service is something customers can buy: give it a bot with prices.
    billing.ensure_bot_for_service(db, key, payload.label, payload.pitch)
    db.commit()
    catalog.invalidate()
    return {"key": key}


@router.put("/services/{key}")
def update_service(key: str, payload: ServiceIn, me: Me = Depends(admin_only),
                   db: Session = Depends(get_db)):
    if key in catalog.BUILTIN:
        raise HTTPException(400, "built-in services cannot be edited")
    svc = db.execute(select(Service).where(Service.key == key)).scalar_one_or_none()
    if not svc:
        raise HTTPException(404, "service not found")
    clash = [x for x in catalog.all_services()
             if x["key"] != key and x["label"].lower() == payload.label.lower()]
    if clash:
        raise HTTPException(409, f"a service called '{payload.label}' already exists")
    for k, v in payload.model_dump().items():
        setattr(svc, k, v)
    db.commit()
    catalog.invalidate()
    return {"ok": True}


@router.delete("/services/{key}")
def delete_service(key: str, me: Me = Depends(admin_only), db: Session = Depends(get_db)):
    if key in catalog.BUILTIN:
        raise HTTPException(400, "built-in services cannot be deleted")
    svc = db.execute(select(Service).where(Service.key == key)).scalar_one_or_none()
    if not svc:
        raise HTTPException(404, "service not found")
    names = db.execute(
        select(Campaign.name).where(Campaign.service_category == key)
    ).scalars().all()
    if names:
        raise HTTPException(
            409, "still used by: " + ", ".join(names) + " - change or delete those campaigns first")
    db.delete(svc)
    db.commit()
    catalog.invalidate()
    return {"ok": True}


@router.get("/campaigns")
def list_campaigns(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.scheduler import upcoming

    campaigns = db.execute(select(Campaign).where(Campaign.owner_id == me.id)
                           .order_by(Campaign.schedule_time)).scalars().all()
    access = billing.access_map(db, me.id)
    next_runs = {u["campaign_id"]: u["at"] for u in upcoming(db, campaigns)}
    # Two grouped queries for all campaigns, not two queries per campaign.
    lead_counts = dict(db.execute(
        select(Lead.campaign_id, func.count(Lead.id)).where(Lead.owner_id == me.id)
        .group_by(Lead.campaign_id)
    ).all())
    latest = select(
        CampaignRun.campaign_id, func.max(CampaignRun.started_at).label("at")
    ).where(CampaignRun.owner_id == me.id).group_by(CampaignRun.campaign_id).subquery()
    last_runs = {r.campaign_id: r for r in db.execute(
        select(CampaignRun).join(latest, (CampaignRun.campaign_id == latest.c.campaign_id)
                                 & (CampaignRun.started_at == latest.c.at))
    ).scalars()}
    out = []
    for c in campaigns:
        last = last_runs.get(c.id)
        leads = lead_counts.get(c.id, 0)
        out.append({
            "id": c.id, "name": c.name, "service_category": c.service_category,
            "enabled": c.enabled, "sources": c.sources, "keywords": c.keywords,
            "locations": c.locations, "countries": c.countries,
            "schedule_time": c.schedule_time, "schedule_days": c.schedule_days,
            "min_score": c.min_score,
            "daily_lead_limit": c.daily_lead_limit,
            "daily_email_limit": c.daily_email_limit,
            "approval_mode": c.approval_mode, "total_leads": leads,
            "rotate_cities": c.rotate_cities,
            "city_coverage": __import__("app.services.cities", fromlist=["x"]).coverage(
                c.countries or [], c.city_history),
            "running": c.id in _running,
            "locked": c.service_category not in access,
            "stopping": pipeline.stop_requested(c.id),
            "next_run": next_runs.get(c.id),
            "last_run": None if not last else {
                "id": last.id, "status": last.status,
                "started_at": as_aware(last.started_at).isoformat() if last.started_at else None,
                "stats": last.stats,
            },
        })
    return out


# Customers' agents stay within a sensible size: every run uses shared search
# quotas, so one account must not be able to exhaust them for everyone.
CUSTOMER_MAX_LEADS_PER_RUN = 100


def _limit_for(me: Me, payload: "CampaignIn") -> None:
    if not me.is_admin and payload.daily_lead_limit > CUSTOMER_MAX_LEADS_PER_RUN:
        raise HTTPException(400, f"Leads per run can be at most {CUSTOMER_MAX_LEADS_PER_RUN}.")


@router.post("/campaigns")
def create_campaign(payload: CampaignIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    _require_bot(db, me, payload.service_category)
    _limit_for(me, payload)
    c = Campaign(owner_id=me.id, **payload.model_dump())
    db.add(c)
    db.commit()
    from app.scheduler import reschedule
    reschedule()
    return {"id": c.id}


@router.put("/campaigns/{campaign_id}")
def update_campaign(campaign_id: int, payload: CampaignIn, me: Me = Depends(current_user),
                    db: Session = Depends(get_db)):
    c = _owned(db, Campaign, campaign_id, me, "campaign")
    if payload.service_category != c.service_category:
        _require_bot(db, me, payload.service_category)
    _limit_for(me, payload)
    for k, v in payload.model_dump().items():
        setattr(c, k, v)
    db.commit()
    from app.scheduler import reschedule
    reschedule()
    return {"ok": True}


@router.delete("/campaigns/{campaign_id}")
def delete_campaign(campaign_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    c = _owned(db, Campaign, campaign_id, me, "campaign")
    # Leads outlive the campaign that found them: detach rather than lose them.
    # Lead.campaign_id has no ON DELETE rule, so Postgres refuses the delete
    # while any lead still points here.
    db.execute(update(Lead).where(Lead.campaign_id == c.id).values(campaign_id=None))
    db.delete(c)
    db.commit()
    from app.scheduler import reschedule
    reschedule()
    return {"ok": True}


def _run_campaign_bg(campaign_id: int, trigger: str) -> None:
    try:
        with session_scope() as db:
            campaign = db.get(Campaign, campaign_id)
            if campaign:
                pipeline.run_campaign(db, campaign, trigger=trigger)
    finally:
        with _running_lock:
            _running.discard(campaign_id)


@router.post("/campaigns/{campaign_id}/stop")
def stop_campaign(campaign_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """Stop a running campaign, or reset one whose run is no longer alive."""
    _owned(db, Campaign, campaign_id, me, "campaign")
    with _running_lock:
        live = campaign_id in _running
    if live:
        pipeline.request_stop(campaign_id)
        return {"stopping": True, "reset": 0}
    # Nothing in this process is executing it: the row was left "running" by a
    # restart or crash. Closing it is the reset.
    closed = pipeline.close_dead_runs(
        db, "Reset by you - this run was no longer active.", campaign_id)
    return {"stopping": False, "reset": closed}


@router.post("/campaigns/{campaign_id}/run")
def run_campaign_now(campaign_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    campaign = _owned(db, Campaign, campaign_id, me, "campaign")
    _require_bot(db, me, campaign.service_category)
    with _running_lock:
        if campaign_id in _running:
            raise HTTPException(409, "this campaign is already running")
        _running.add(campaign_id)
    threading.Thread(
        target=_run_campaign_bg, args=(campaign_id, "manual"), daemon=True
    ).start()
    return {"started": True, "campaign_id": campaign_id}


@router.get("/campaigns/{campaign_id}/runs")
def campaign_runs(campaign_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    _owned(db, Campaign, campaign_id, me, "campaign")
    runs = db.execute(
        select(CampaignRun).where(CampaignRun.campaign_id == campaign_id)
        .order_by(CampaignRun.started_at.desc()).limit(20)
    ).scalars().all()
    return [{
        "id": r.id, "status": r.status, "trigger": r.trigger,
        "started_at": r.started_at.isoformat() if r.started_at else None,
        "finished_at": r.finished_at.isoformat() if r.finished_at else None,
        "stats": r.stats, "error": r.error,
    } for r in runs]


@router.get("/workflow")
def workflow_state(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """Live state of the agent graph for the main dashboard.

    Returns the most recent run - active if one is in flight, otherwise the
    last finished one, so the canvas always has something real to show.
    `upcoming` says which campaign starts next and when, for the countdown.
    """
    from app.scheduler import upcoming
    from app.services.pipeline import PIPELINE_STAGES

    schedule = {
        "upcoming": upcoming(db, db.execute(
            select(Campaign).where(Campaign.owner_id == me.id, Campaign.enabled.is_(True))
        ).scalars().all()),
        "scheduler_enabled": settings.scheduler_enabled,
        "timezone": settings.scheduler_timezone,
    }

    # A running run wins; otherwise the newest. One query, campaign included.
    found = db.execute(
        select(CampaignRun, Campaign).outerjoin(Campaign, Campaign.id == CampaignRun.campaign_id)
        .where(CampaignRun.owner_id == me.id)
        .order_by((CampaignRun.status == "running").desc(), CampaignRun.started_at.desc())
        .limit(1)
    ).first()
    run, campaign = found if found else (None, None)

    idle_nodes = [
        {"key": k, "label": lab, "note": note, "icon": ic, "engine": eng,
         "status": "idle", "passed": 0, "dropped": 0, "detail": "",
         "drop_detail": "", "ms": 0}
        for k, lab, note, ic, eng in PIPELINE_STAGES
    ]
    if run is None:
        return {
            "run": None, "nodes": idle_nodes, "log": [],
            "campaigns_ready": db.execute(select(func.count(Campaign.id))
                                          .where(Campaign.owner_id == me.id)).scalar_one(),
            **schedule,
        }

    state = run.stage_state or {}
    nodes = state.get("nodes") or idle_nodes
    return {
        "active": state.get("active"),
        "run": {
            "id": run.id, "status": run.status, "trigger": run.trigger,
            "campaign": campaign.name if campaign else "—",
            "campaign_id": run.campaign_id,
            "service": campaign.service_category if campaign else "",
            # Explicit offset: SQLite returns naive values, which a browser
            # would read as local time and shift the elapsed clock by hours.
            "started_at": as_aware(run.started_at).isoformat() if run.started_at else None,
            "finished_at": as_aware(run.finished_at).isoformat() if run.finished_at else None,
            "current_item": run.current_item,
            "stats": run.stats or {},
            "error": run.error,
            "stopping": run.status == "running" and pipeline.stop_requested(run.campaign_id),
        },
        "nodes": nodes,
        "log": (run.log or [])[-60:],
        **schedule,
    }


@router.get("/runs/{run_id}")
def run_detail(run_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    r = _owned(db, CampaignRun, run_id, me, "run")
    queries = db.execute(
        select(SearchQuery).where(SearchQuery.run_id == run_id)
    ).scalars().all()
    return {
        "id": r.id, "campaign_id": r.campaign_id, "status": r.status,
        "started_at": r.started_at.isoformat() if r.started_at else None,
        "finished_at": r.finished_at.isoformat() if r.finished_at else None,
        "stats": r.stats, "log": r.log, "error": r.error,
        "nodes": (r.stage_state or {}).get("nodes") or [],
        "current_item": r.current_item,
        "queries": [{"source": q.source, "query": q.query, "results": q.results_count} for q in queries],
    }


@router.get("/runs")
def recent_runs(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.execute(
        select(CampaignRun, Campaign.name)
        .join(Campaign, Campaign.id == CampaignRun.campaign_id)
        .where(CampaignRun.owner_id == me.id)
        .order_by(CampaignRun.started_at.desc()).limit(15)
    ).all()
    return [{
        "id": r.id, "campaign": name, "status": r.status, "trigger": r.trigger,
        "started_at": r.started_at.isoformat() if r.started_at else None,
        "stats": r.stats,
    } for r, name in rows]


# ---------------------------------------------------------------------------
# Leads
# ---------------------------------------------------------------------------


def _lead_row(lead: Lead, company: Company, contact: Contact | None) -> dict:
    return {
        "id": lead.id, "company_id": company.id, "company": company.name,
        "country": company.country, "city": company.city,
        "website": company.website, "industry": company.industry,
        "phone": company.phone,
        "contact_name": contact.full_name if contact else None,
        "email": contact.email if contact else None,
        "service_category": lead.service_category, "intent_level": lead.intent_level,
        "score": lead.score, "stage": lead.stage, "status": lead.status,
        "source_kind": lead.source_kind, "intent_reason": lead.intent_reason,
        "email_policy": company.email_policy,
        "created_at": lead.created_at.isoformat() if lead.created_at else None,
        "last_contact_at": lead.last_contact_at.isoformat() if lead.last_contact_at else None,
    }


@router.get("/leads")
def list_leads(
    service: str = "", source: str = "", country: str = "", intent: str = "",
    status: str = "", stage: str = "", min_score: int = 0, q: str = "",
    date_from: str = "", date_to: str = "", has_email: str = "",
    sort: str = "score", page: int = 1, page_size: int = 50,
    me: Me = Depends(current_user), db: Session = Depends(get_db),
):
    page_size = max(1, min(page_size, 200))
    stmt = select(Lead, Company, Contact).join(Company, Company.id == Lead.company_id).outerjoin(
        Contact, Contact.id == Lead.contact_id
    ).where(Lead.owner_id == me.id)
    if service:
        stmt = stmt.where(Lead.service_category == service)
    if source:
        stmt = stmt.where(Lead.source_kind == source)
    if country:
        stmt = stmt.where(Company.country == country.upper())
    if intent:
        stmt = stmt.where(Lead.intent_level == intent)
    if status:
        stmt = stmt.where(Lead.status == status)
    if stage:
        stmt = stmt.where(Lead.stage == stage)
    if min_score:
        stmt = stmt.where(Lead.score >= min_score)
    if has_email == "yes":
        stmt = stmt.where(Lead.contact_id.isnot(None))
    elif has_email == "no":
        stmt = stmt.where(Lead.contact_id.is_(None))
    if q:
        like = f"%{q.lower()}%"
        stmt = stmt.where(func.lower(Company.name).like(like))
    if date_from:
        stmt = stmt.where(func.date(Lead.created_at) >= date_from)
    if date_to:
        stmt = stmt.where(func.date(Lead.created_at) <= date_to)

    total = db.execute(
        select(func.count()).select_from(stmt.subquery())
    ).scalar_one()

    order = {
        "score": Lead.score.desc(), "created": Lead.created_at.desc(),
        "company": Company.name.asc(), "oldest": Lead.created_at.asc(),
    }.get(sort, Lead.score.desc())
    stmt = stmt.order_by(order, Lead.id.desc()).limit(page_size).offset((page - 1) * page_size)

    rows = db.execute(stmt).all()
    return {
        "total": total, "page": page, "page_size": page_size,
        "leads": [_lead_row(l, c, ct) for l, c, ct in rows],
    }


@router.get("/leads/{lead_id}")
def lead_detail(lead_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    row = db.execute(
        select(Lead, Company, Contact).join(Company, Company.id == Lead.company_id)
        .outerjoin(Contact, Contact.id == Lead.contact_id)
        .where(Lead.id == lead_id, Lead.owner_id == me.id)
    ).first()
    if not row:
        raise HTTPException(404, "lead not found")
    lead, company, contact = row

    # Evidence and the pages it came from together (this used to load every
    # source in the database, however many there were).
    pairs = db.execute(
        select(Signal, Source).outerjoin(Source, Source.id == Signal.source_id)
        .where(Signal.company_id == company.id).order_by(Signal.observed_at.desc())
    ).all()
    signals = [s for s, _ in pairs]
    sources = {src.id: src for _, src in pairs if src is not None}
    audit = db.execute(
        select(WebsiteAudit).where(WebsiteAudit.company_id == company.id)
        .order_by(WebsiteAudit.audited_at.desc()).limit(1)
    ).scalar_one_or_none()
    analyses = db.execute(
        select(AIAnalysis).where(AIAnalysis.lead_id == lead_id).order_by(AIAnalysis.created_at)
    ).scalars().all()
    messages = db.execute(
        select(EmailMessage).where(EmailMessage.lead_id == lead_id, EmailMessage.owner_id == me.id)
        .order_by(EmailMessage.created_at.desc())
    ).scalars().all()
    contacts = db.execute(
        select(Contact).where(Contact.company_id == company.id)
    ).scalars().all()

    policy, policy_reason = policy_for_country(company.country)
    return {
        **_lead_row(lead, company, contact),
        "address": company.address,
        "notes": lead.notes,
        "score_breakdown": lead.score_breakdown,
        "policy_reason": policy_reason,
        "contacts": [
            {"id": c.id, "email": c.email, "name": c.full_name, "title": c.title,
             "status": c.email_status} for c in contacts
        ],
        "signals": [{
            "id": s.id, "type": s.signal_type, "evidence_class": s.evidence_class,
            "quote": s.quote, "detected_by": s.detected_by,
            "observed_at": s.observed_at.isoformat() if s.observed_at else None,
            "source_url": sources[s.source_id].url if s.source_id in sources else "",
            "source_kind": sources[s.source_id].kind if s.source_id in sources else "",
        } for s in signals],
        "audit": None if not audit else {
            "url": audit.url, "has_site": audit.has_site, "has_ssl": audit.has_ssl,
            "is_mobile_friendly": audit.is_mobile_friendly, "has_booking": audit.has_booking,
            "has_ecommerce": audit.has_ecommerce, "has_contact_form": audit.has_contact_form,
            "copyright_year": audit.copyright_year, "tech": audit.tech,
            "ttfb_ms": audit.ttfb_ms, "notes": audit.notes,
            "audited_at": audit.audited_at.isoformat() if audit.audited_at else None,
        },
        "ai_analyses": [{
            "agent": a.agent, "model": a.model, "output": a.output,
            "cost_usd": a.cost_usd, "cited": a.output.get("cited_signal_ids", []),
            "created_at": a.created_at.isoformat() if a.created_at else None,
        } for a in analyses],
        "messages": [{
            "id": m.id, "subject": m.subject, "body": m.body, "status": m.status,
            "to_email": m.to_email, "cited_signal_ids": m.cited_signal_ids,
            "failure_reason": m.failure_reason,
            "created_at": m.created_at.isoformat() if m.created_at else None,
            "sent_at": m.sent_at.isoformat() if m.sent_at else None,
        } for m in messages],
    }


class LeadPatch(BaseModel):
    status: str | None = None
    notes: str | None = None


@router.patch("/leads/{lead_id}")
def patch_lead(lead_id: int, payload: LeadPatch, me: Me = Depends(current_user),
               db: Session = Depends(get_db)):
    lead = _owned(db, Lead, lead_id, me, "lead")
    if payload.status is not None:
        lead.status = payload.status
    if payload.notes is not None:
        lead.notes = payload.notes
    db.commit()
    return {"ok": True}


@router.post("/leads/{lead_id}/rescore")
def rescore_lead(lead_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    lead = _owned(db, Lead, lead_id, me, "lead")
    company = db.get(Company, lead.company_id)
    contact = db.get(Contact, lead.contact_id) if lead.contact_id else None
    signals = company_signals(db, company.id)
    result = scoring.score_lead(
        company=company, signals=signals, contact=contact,
        service_category=lead.service_category,
    )
    lead.score, lead.intent_level = result.score, result.intent_level
    lead.score_breakdown, lead.score_version = result.breakdown, scoring.SCORE_VERSION
    db.commit()
    return {"score": result.score, "intent_level": result.intent_level,
            "breakdown": result.breakdown}


@router.post("/leads/{lead_id}/draft")
def draft_lead(lead_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    lead = _owned(db, Lead, lead_id, me, "lead")
    company = db.get(Company, lead.company_id)
    contact = db.get(Contact, lead.contact_id) if lead.contact_id else None
    if not contact:
        raise HTTPException(400, "this lead has no email address")
    signals = company_signals(db, company.id)
    if not signals:
        raise HTTPException(400, "no evidence on file — refusing to write outreach")

    draft = ai.draft_email(company, signals, contact, lead.service_category,
                           mailer.sender_profile(db, me.id))
    usage = draft.pop("_usage", {})
    db.add(AIAnalysis(
        lead_id=lead.id, agent="personalize", model=usage.get("model", ""),
        input_signal_ids=[s.id for s in signals],
        output={k: v for k, v in draft.items() if not k.startswith("_")},
        tokens_in=usage.get("tokens_in", 0), tokens_out=usage.get("tokens_out", 0),
        cost_usd=usage.get("cost_usd", 0.0),
    ))
    message = EmailMessage(
        owner_id=me.id, lead_id=lead.id, contact_id=contact.id, to_email=contact.email,
        subject=draft["subject"], body=draft["body"],
        cited_signal_ids=draft.get("cited_signal_ids", []),
        unsubscribe_token=mailer.new_unsubscribe_token(),
        status="pending_approval", scheduled_for=utcnow(),
    )
    db.add(message)
    lead.stage = "pending_approval"
    db.commit()
    return {"message_id": message.id, "subject": message.subject, "body": message.body,
            "used_fallback": draft.get("_fallback", False)}


# ---------------------------------------------------------------------------
# Outbox
# ---------------------------------------------------------------------------


@router.get("/outbox")
def outbox(status: str = "pending_approval", me: Me = Depends(current_user),
           db: Session = Depends(get_db)):
    rows = db.execute(
        select(EmailMessage, Lead, Company)
        .join(Lead, Lead.id == EmailMessage.lead_id)
        .join(Company, Company.id == Lead.company_id)
        .where(EmailMessage.owner_id == me.id, EmailMessage.status == status)
        .order_by(Lead.score.desc(), EmailMessage.created_at.asc())
        .limit(200)
    ).all()
    evidence = signals_by_company(db, (company.id for _, _, company in rows))
    out = []
    for m, lead, company in rows:
        signals = evidence.get(company.id, [])
        cited = set(m.cited_signal_ids or [])
        out.append({
            "id": m.id, "lead_id": lead.id, "company": company.name,
            "country": company.country, "website": company.website,
            "to_email": m.to_email, "subject": m.subject, "body": m.body,
            "score": lead.score, "intent_level": lead.intent_level,
            "service_category": lead.service_category, "status": m.status,
            "failure_reason": m.failure_reason,
            "created_at": m.created_at.isoformat() if m.created_at else None,
            "evidence": [{
                "id": s.id, "type": s.signal_type, "quote": s.quote,
                "evidence_class": s.evidence_class, "cited": s.id in cited,
            } for s in signals],
        })
    return out


class MessagePatch(BaseModel):
    subject: str | None = None
    body: str | None = None


@router.patch("/outbox/{message_id}")
def edit_message(message_id: int, payload: MessagePatch, me: Me = Depends(current_user),
                 db: Session = Depends(get_db)):
    m = _owned(db, EmailMessage, message_id, me, "message")
    if m.status not in ("pending_approval", "draft", "queued"):
        raise HTTPException(400, f"cannot edit a message that is '{m.status}'")
    if payload.subject is not None:
        m.subject = payload.subject
    if payload.body is not None:
        m.body = payload.body
    db.commit()
    return {"ok": True}


@router.post("/outbox/{message_id}/approve")
def approve_message(message_id: int, send_now: bool = True, me: Me = Depends(current_user),
                    db: Session = Depends(get_db)):
    m = _owned(db, EmailMessage, message_id, me, "message")
    lead = db.get(Lead, m.lead_id)
    company = db.get(Company, lead.company_id)
    contact = db.get(Contact, m.contact_id) if m.contact_id else None

    gate = mailer.check_can_send(db, lead, company, contact, exclude_message_id=m.id)
    if not gate.ok:
        raise HTTPException(400, f"cannot send: {gate.reason}")

    m.status = "queued"
    m.approved_at = utcnow()
    m.approved_by = "admin"
    lead.stage = "queued"
    db.commit()

    if send_now and settings.send_on_server:
        sent, detail = mailer.send_message(db, m)
        db.commit()
        # Still queued means the server could not reach the mail server: the
        # approval stands and the email goes out on the next attempt.
        return {"ok": sent or m.status == "queued", "detail": detail, "status": m.status}
    return {"ok": True, "detail": "queued", "status": m.status}


@router.post("/outbox/{message_id}/reject")
def reject_message(message_id: int, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    m = _owned(db, EmailMessage, message_id, me, "message")
    m.status = "rejected"
    lead = db.get(Lead, m.lead_id)
    if lead:
        lead.stage = "rejected"
        lead.status = "rejected_by_admin"
    db.commit()
    return {"ok": True}


@router.post("/outbox/dispatch")
def dispatch_now(limit: int = 25, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    if not settings.send_on_server:
        return {"sent": 0, "failed": 0, "skipped": 0,
                "details": ["Emails are sent from your computer (run.py --send-only), not this server."]}
    return pipeline.dispatch_queue(db, limit=min(limit, 50), owner_id=me.id)


@router.get("/outbox/summary")
def outbox_summary(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.services import sender
    return {**sender.outbox_summary(db, me.id), "send_on_server": settings.send_on_server}


@router.post("/outbox/retry-failed")
def retry_failed(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """Put failed sends back in the queue. Every check runs again at send time."""
    msgs = db.execute(select(EmailMessage).where(EmailMessage.owner_id == me.id,
                                                 EmailMessage.status == "failed")).scalars().all()
    for m in msgs:
        m.status = "queued"
        m.failure_reason = None
        m.scheduled_for = utcnow()
    db.commit()
    return {"requeued": len(msgs)}


# ---------------------------------------------------------------------------
# Suppressions & unsubscribe
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# WhatsApp, alerts, inbox
# ---------------------------------------------------------------------------


@router.get("/linkedin")
def linkedin_queue(only_unemailable: bool = False, me: Me = Depends(current_user),
                   db: Session = Depends(get_db)):
    return {
        "stats": linkedin.stats(db, me.id),
        "leads": linkedin.queue(db, owner_id=me.id, only_unemailable=only_unemailable),
    }


@router.get("/whatsapp")
def whatsapp_queue(status: str = "pending", me: Me = Depends(current_user),
                   db: Session = Depends(get_db)):
    rows = db.execute(
        select(WhatsAppMessage, Lead, Company)
        .join(Lead, Lead.id == WhatsAppMessage.lead_id)
        .join(Company, Company.id == Lead.company_id)
        .where(WhatsAppMessage.owner_id == me.id, WhatsAppMessage.status == status)
        .order_by(Lead.score.desc(), WhatsAppMessage.created_at.asc())
        .limit(200)
    ).all()
    return [{
        "id": m.id, "lead_id": lead.id, "company": company.name,
        "country": company.country, "phone": m.phone_e164, "body": m.body,
        "score": lead.score, "status": m.status,
        "service_category": lead.service_category,
        "link": whatsapp.wa_link(m.phone_e164, m.body),
        "created_at": m.created_at.isoformat() if m.created_at else None,
    } for m, lead, company in rows]


@router.post("/whatsapp/{message_id}/{action}")
def whatsapp_mark(message_id: int, action: str, me: Me = Depends(current_user),
                  db: Session = Depends(get_db)):
    if action not in ("opened", "sent", "skipped"):
        raise HTTPException(400, "action must be opened, sent or skipped")
    if not whatsapp.mark(db, message_id, action, owner_id=me.id):
        raise HTTPException(404, "message not found")
    db.commit()
    return {"ok": True}


@router.post("/notify/test")
def notify_test(me: Me = Depends(admin_only)):
    if not notify.ready():
        raise HTTPException(400, "TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID are not set")
    ok, detail = notify.send(
        "<b>✅ Money Maker connected</b>\n\nReply alerts will arrive here."
    )
    if not ok:
        raise HTTPException(502, f"Telegram rejected it: {detail}")
    return {"ok": True}


@router.post("/inbox/poll")
def inbox_poll(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    cfg = tenant.config_for(db, me.id)
    if not cfg.imap_ready:
        raise HTTPException(400, "No inbox connected - add your IMAP details in Settings")
    return inbox.poll(db, cfg)


@router.get("/suppressions")
def list_suppressions(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.execute(select(Suppression).where(Suppression.owner_id == me.id)
                      .order_by(Suppression.created_at.desc()).limit(500)).scalars().all()
    return [{"id": s.id, "scope": s.scope, "value": s.value, "reason": s.reason,
             "created_at": s.created_at.isoformat() if s.created_at else None} for s in rows]


@router.post("/suppressions")
def add_suppression(
    value: str = Body(..., embed=True), scope: str = Body("email", embed=True),
    reason: str = Body("added manually", embed=True), me: Me = Depends(current_user),
    db: Session = Depends(get_db),
):
    if scope not in ("email", "domain"):
        raise HTTPException(400, "scope must be email or domain")
    suppress(db, value, reason, scope, owner_id=me.id)
    db.commit()
    return {"ok": True}


@router.delete("/suppressions/{suppression_id}")
def remove_suppression(suppression_id: int, me: Me = Depends(current_user),
                       db: Session = Depends(get_db)):
    db.execute(delete(Suppression).where(Suppression.id == suppression_id,
                                         Suppression.owner_id == me.id))
    db.commit()
    return {"ok": True}


@router.get("/mailboxes")
def list_mailboxes(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    rows = db.execute(select(Mailbox).where(Mailbox.owner_id == me.id)).scalars().all()
    return [{
        "id": m.id, "address": m.address, "enabled": m.enabled,
        "daily_cap": m.daily_cap, "effective_cap_today": mailer.warmup_cap(m),
        "sent_today": mailer.sent_today(db, m.id),
        "warmup_started_at": m.warmup_started_at.isoformat() if m.warmup_started_at else None,
    } for m in rows]


@router.get("/settings")
def read_settings(me: Me = Depends(admin_only)):
    """Non-secret configuration, plus whether each secret is present."""
    return {
        "company_name": settings.company_name,
        "company_website": settings.company_website,
        "company_postal_address": settings.company_postal_address,
        "sender_name": settings.sender_name,
        "sender_email": settings.sender_email,
        "portfolio_url": settings.portfolio_url,
        "privacy_policy_url": settings.privacy_policy_url,
        "public_base_url": settings.public_base_url,
        "daily_email_cap_global": settings.daily_email_cap_global,
        "per_mailbox_daily_cap": settings.per_mailbox_daily_cap,
        "cooling_off_days": settings.cooling_off_days,
        "min_score_to_email": settings.min_score_to_email,
        "configured": {
            "anthropic_api_key": bool(settings.anthropic_api_key),
            "companies_house_key": bool(settings.companies_house_key),
            "adzuna": bool(settings.adzuna_app_id and settings.adzuna_app_key),
            "smtp": bool(settings.smtp_host),
        },
    }


@router.delete("/leads/purge-evidenceless")
def purge_evidenceless(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """Delete this user's leads whose company has no evidence on file.

    A lead with no signal is a name, not a prospect: nothing can be said to
    them truthfully, so it can never be emailed. Companies are shared between
    customers, so a company is only removed once no lead points at it.
    """
    company_ids = set(db.execute(select(Signal.company_id)).scalars().all())
    leads = db.execute(select(Lead).where(Lead.owner_id == me.id)).scalars().all()
    removed, kept, orphans = 0, 0, set()
    for lead in leads:
        if lead.company_id in company_ids:
            kept += 1
            continue
        orphans.add(lead.company_id)
        db.delete(lead)
        removed += 1
    db.flush()
    for cid in orphans:
        if not db.execute(select(Lead.id).where(Lead.company_id == cid).limit(1)).first():
            company = db.get(Company, cid)
            if company:
                db.delete(company)
    db.commit()
    return {"removed": removed, "kept": kept}


@router.post("/legacy/import")
def legacy_import(me: Me = Depends(admin_only), db: Session = Depends(get_db)):
    """Opt-in import of the previous version's leads, as history only."""
    from app.seed import migrate_legacy

    return {"imported": migrate_legacy(db)}


@router.get("/sources")
def source_guide(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """What each source is for, whether it is ready, and which key to get next."""
    from app.seed import DEFAULT_CAMPAIGNS, FREE_KEY_PRIORITY

    by_service = {
        c["service_category"]: c["sources"] for c in DEFAULT_CAMPAIGNS
    }
    return {
        "sources": available_sources(),
        "recommended": by_service,
        "free_keys": [
            {**k, "configured": _key_configured(k["name"])} for k in FREE_KEY_PRIORITY
        ] if me.is_admin else [],
    }


def _key_configured(name: str) -> bool:
    return {
        "Serper": bool(settings.serper_api_key),
        "Adzuna": bool(settings.adzuna_app_id and settings.adzuna_app_key),
        "Companies House": bool(settings.companies_house_key),
        "SAM.gov": bool(getattr(settings, "sam_gov_api_key", "")),
    }.get(name, False)


@router.post("/sources/apply-recommended")
def apply_recommended(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.seed import apply_recommended_sources
    from app.scheduler import reschedule

    changed = apply_recommended_sources(db, me.id)
    reschedule()
    return {"ok": True, "campaigns_updated": changed}


@router.get("/export/leads.xlsx")
def export_leads(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = db.execute(
        select(Lead, Company, Contact)
        .join(Company, Company.id == Lead.company_id)
        .outerjoin(Contact, Contact.id == Lead.contact_id)
        .where(Lead.owner_id == me.id)
        .order_by(Lead.score.desc())
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = ["Score", "Intent", "Company", "Country", "City", "Industry", "Website",
               "Contact", "Email", "Phone", "Service", "Source", "Stage", "Status",
               "Why this lead", "Created"]
    fill, font = PatternFill("solid", fgColor="0E7490"), Font(color="FFFFFF", bold=True)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.alignment = fill, font, Alignment(horizontal="left")
    for r, (lead, company, contact) in enumerate(rows, 2):
        for i, v in enumerate([
            lead.score, lead.intent_level, company.name, company.country, company.city,
            company.industry, company.website, contact.full_name if contact else "",
            contact.email if contact else "", company.phone, lead.service_category,
            lead.source_kind, lead.stage, lead.status, lead.intent_reason,
            lead.created_at.strftime("%Y-%m-%d %H:%M") if lead.created_at else "",
        ], 1):
            ws.cell(row=r, column=i, value=v)
    for i, w in enumerate([7, 9, 30, 8, 16, 22, 34, 20, 30, 18, 18, 16, 14, 16, 60, 18], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"moneymaker_leads_{dt.date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


# ---------------------------------------------------------------------------
# Account: profile, sending identity and mailbox
# ---------------------------------------------------------------------------


class AccountIn(BaseModel):
    name: str | None = Field(None, max_length=160)
    company: str | None = Field(None, max_length=160)
    sender_name: str | None = Field(None, max_length=160)
    sender_email: str | None = Field(None, max_length=320)
    company_name: str | None = Field(None, max_length=160)
    company_website: str | None = Field(None, max_length=512)
    postal_address: str | None = Field(None, max_length=600)
    privacy_url: str | None = Field(None, max_length=512)
    smtp_host: str | None = Field(None, max_length=255)
    smtp_port: int | None = Field(None, ge=1, le=65535)
    smtp_use_ssl: bool | None = None
    smtp_user: str | None = Field(None, max_length=320)
    smtp_pass: str | None = Field(None, max_length=500)  # blank keeps the saved one
    imap_host: str | None = Field(None, max_length=255)
    imap_port: int | None = Field(None, ge=1, le=65535)
    imap_user: str | None = Field(None, max_length=320)
    imap_pass: str | None = Field(None, max_length=500)  # blank keeps the saved one
    imap_folder: str | None = Field(None, max_length=120)
    daily_cap: int | None = Field(None, ge=1, le=500)
    timezone: str | None = Field(None, max_length=64)

    @field_validator("sender_email")
    @classmethod
    def _email(cls, v: str | None) -> str | None:
        v = (v or "").strip()
        if not v:
            return None
        from email_validator import EmailNotValidError, validate_email
        try:
            validate_email(v, check_deliverability=False)
        except EmailNotValidError as exc:
            raise ValueError(f"sender email is not valid: {exc}") from None
        return v.lower()

    @field_validator("timezone")
    @classmethod
    def _tz(cls, v: str | None) -> str | None:
        if not v:
            return None
        from zoneinfo import ZoneInfo
        try:
            ZoneInfo(v)
        except Exception:  # noqa: BLE001
            raise ValueError(f"unknown timezone '{v}'") from None
        return v


ACCOUNT_FIELDS = ("sender_name", "sender_email", "company_name", "company_website",
                  "postal_address", "privacy_url", "smtp_host", "smtp_port", "smtp_use_ssl",
                  "smtp_user", "imap_host", "imap_port", "imap_user", "imap_folder",
                  "daily_cap", "timezone")


@router.get("/account")
def read_account(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import AccountSettings, User

    user = db.get(User, me.id)
    row = db.get(AccountSettings, me.id)
    cfg = tenant.config_for(db, me.id)
    return {
        "name": user.name, "email": user.email, "company": user.company, "role": user.role,
        "created_at": user.created_at.isoformat() if user.created_at else None,
        "saved": {f: getattr(row, f, None) if row else None for f in ACCOUNT_FIELDS},
        "has_smtp_pass": bool(row and row.smtp_pass_enc) or bool(cfg.smtp_pass),
        "has_imap_pass": bool(row and row.imap_pass_enc) or bool(cfg.imap_pass),
        # What is actually used, after the admin's fall-back to the server settings.
        "effective": {
            "sender_name": cfg.sender_name, "sender_email": cfg.sender_email,
            "company_name": cfg.company_name, "company_website": cfg.company_website,
            "postal_address": cfg.postal_address, "privacy_url": cfg.privacy_url,
            "smtp_host": cfg.smtp_host, "smtp_port": cfg.smtp_port, "smtp_use_ssl": cfg.smtp_use_ssl,
            "smtp_user": cfg.smtp_user, "imap_host": cfg.imap_host, "imap_port": cfg.imap_port,
            "imap_user": cfg.imap_user, "imap_folder": cfg.imap_folder,
            "daily_cap": cfg.daily_cap, "timezone": cfg.timezone,
        },
        "ready": {"smtp": cfg.smtp_ready, "imap": cfg.imap_ready, "sender": cfg.sender_ready},
    }


@router.put("/account")
def update_account(payload: AccountIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import AccountSettings, User

    user = db.get(User, me.id)
    if payload.name is not None:
        user.name = payload.name.strip() or user.name
    if payload.company is not None:
        user.company = payload.company.strip() or None
    row = db.get(AccountSettings, me.id)
    if row is None:
        row = AccountSettings(owner_id=me.id)
        db.add(row)
    data = payload.model_dump()
    for f in ACCOUNT_FIELDS:
        value = data.get(f)
        if isinstance(value, str):
            value = value.strip() or None
        setattr(row, f, value)
    if payload.smtp_pass:
        row.smtp_pass_enc = tenant.encrypt(payload.smtp_pass)
    if payload.imap_pass:
        row.imap_pass_enc = tenant.encrypt(payload.imap_pass)
    db.commit()
    from app.scheduler import reschedule
    reschedule()  # a timezone change moves this user's agents
    return {"ok": True}


@router.post("/account/test-mail")
def test_mailbox(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    """Log in to the saved SMTP and IMAP servers, so mistakes show up now, not at send time."""
    import imaplib
    import smtplib
    import ssl

    cfg = tenant.config_for(db, me.id)
    out = {"smtp": {"ok": False, "detail": "not set"}, "imap": {"ok": False, "detail": "not set"}}
    if cfg.smtp_host:
        try:
            ctx = ssl.create_default_context()
            if cfg.smtp_use_ssl or cfg.smtp_port == 465:
                conn = smtplib.SMTP_SSL(cfg.smtp_host, cfg.smtp_port, timeout=20, context=ctx)
            else:
                conn = smtplib.SMTP(cfg.smtp_host, cfg.smtp_port, timeout=20)
                conn.ehlo()
                conn.starttls(context=ctx)
            with conn:
                if cfg.smtp_user:
                    conn.login(cfg.smtp_user, cfg.smtp_pass)
            out["smtp"] = {"ok": True, "detail": f"Logged in as {cfg.smtp_user or '(no login)'}"}
        except Exception as exc:  # noqa: BLE001
            out["smtp"] = {"ok": False, "detail": f"{type(exc).__name__}: {str(exc)[:200]}"}
    if cfg.imap_ready:
        try:
            with imaplib.IMAP4_SSL(cfg.imap_host, cfg.imap_port, timeout=20) as box:
                box.login(cfg.imap_user or cfg.smtp_user, cfg.imap_pass or cfg.smtp_pass)
            out["imap"] = {"ok": True, "detail": f"Logged in as {cfg.imap_user or cfg.smtp_user}"}
        except Exception as exc:  # noqa: BLE001
            out["imap"] = {"ok": False, "detail": f"{type(exc).__name__}: {str(exc)[:200]}"}
    return out


class PasswordIn(BaseModel):
    current: str
    new: str = Field(min_length=8, max_length=200)


@router.post("/account/password")
def change_password(payload: PasswordIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import User

    user = db.get(User, me.id)
    if me.is_admin:
        raise HTTPException(400, "The admin password is set by ADMIN_PASSWORD on the server.")
    if not auth.verify_password(user.password_hash, payload.current):
        raise HTTPException(400, "Your current password is not correct.")
    user.password_hash = auth.hash_password(payload.new)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------------------
# Billing: trial, checkout, verification
# ---------------------------------------------------------------------------


def _unlock(db: Session, user_id: int, bot_key: str) -> None:
    """After access is granted: a ready-made agent for that bot, scheduled."""
    from app.scheduler import reschedule
    from app.seed import starter_campaign

    starter_campaign(db, user_id, bot_key)
    db.commit()
    reschedule()


@router.get("/billing")
def billing_overview(me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import Payment, Subscription, User

    user = db.get(User, me.id)
    subs = db.execute(select(Subscription).where(Subscription.user_id == me.id)
                      .order_by(Subscription.created_at.desc()).limit(100)).scalars().all()
    pays = db.execute(select(Payment).where(Payment.user_id == me.id)
                      .order_by(Payment.created_at.desc()).limit(100)).scalars().all()
    return {
        "access": billing.access_map(db, me.id),
        "bots": [billing.bot_json(b) for b in billing.catalog(db)],
        "trial": {"used": bool(user.trial_started_at), "bot": user.trial_bot,
                  "hours": settings.trial_hours},
        "razorpay_ready": billing.razorpay_ready(),
        "subscriptions": [{
            "id": s.id, "bot": s.bot_key, "kind": s.kind, "period": s.period,
            "currency": s.currency, "amount": s.amount,
            "starts_at": as_aware(s.starts_at).isoformat(), "expires_at": as_aware(s.expires_at).isoformat(),
            "revoked": s.revoked_at is not None, "note": s.note,
        } for s in subs],
        "payments": [{
            "id": p.id, "bot": p.bot_key, "period": p.period, "currency": p.currency,
            "amount": p.amount, "status": p.status, "order_id": p.order_id,
            "payment_id": p.provider_payment_id,
            "created_at": as_aware(p.created_at).isoformat() if p.created_at else None,
            "paid_at": as_aware(p.paid_at).isoformat() if p.paid_at else None,
        } for p in pays],
    }


class TrialIn(BaseModel):
    bot: str


@router.post("/billing/trial")
def start_trial(payload: TrialIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import User

    if me.is_admin:
        raise HTTPException(400, "The admin already has every bot.")
    user = db.get(User, me.id)
    try:
        sub = billing.start_trial(db, user, payload.bot)
    except billing.TrialError as exc:
        raise HTTPException(400, str(exc)) from None
    db.commit()
    _unlock(db, me.id, payload.bot)
    return {"ok": True, "expires_at": as_aware(sub.expires_at).isoformat()}


class CheckoutIn(BaseModel):
    bot: str
    period: str = "month"
    currency: str = "INR"


@router.post("/billing/checkout")
def checkout(payload: CheckoutIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import User

    if me.is_admin:
        raise HTTPException(400, "The admin already has every bot.")
    try:
        order = billing.create_order(db, db.get(User, me.id), payload.bot,
                                     payload.period, payload.currency.upper())
    except billing.PaymentError as exc:
        raise HTTPException(400, str(exc)) from None
    db.commit()
    return order


class VerifyIn(BaseModel):
    order_id: str
    payment_id: str
    signature: str


@router.post("/billing/verify")
def verify_payment(payload: VerifyIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import Payment

    payment = db.execute(select(Payment).where(Payment.order_id == payload.order_id)).scalar_one_or_none()
    if payment is None or payment.user_id != me.id:
        raise HTTPException(404, "payment not found")
    if not billing.checkout_signature_ok(payload.order_id, payload.payment_id, payload.signature):
        raise HTTPException(400, "The payment could not be verified. If money was taken, "
                                 "it will be confirmed automatically within a few minutes.")
    sub = billing.fulfil(db, payload.order_id, payload.payment_id)
    db.commit()
    if sub is None:
        raise HTTPException(400, "This payment could not be applied. Please contact support.")
    _unlock(db, me.id, payment.bot_key)
    return {"ok": True, "bot": payment.bot_key, "expires_at": as_aware(sub.expires_at).isoformat()}


class FailedIn(BaseModel):
    order_id: str
    reason: str = ""


@router.post("/billing/failed")
def payment_failed(payload: FailedIn, me: Me = Depends(current_user), db: Session = Depends(get_db)):
    from app.models import Payment

    payment = db.execute(select(Payment).where(Payment.order_id == payload.order_id)).scalar_one_or_none()
    if payment and payment.user_id == me.id and payment.status == "created":
        payment.status, payment.error = "failed", (payload.reason or "")[:500]
        db.commit()
    return {"ok": True}


# --- public, unauthenticated ----------------------------------------------

public_router = APIRouter()


@public_router.api_route("/u/{token}", methods=["GET", "POST"], response_class=HTMLResponse)
def unsubscribe(token: str, db: Session = Depends(get_db)):
    """One-click unsubscribe. POST is what RFC 8058 clients call."""
    message = db.execute(
        select(EmailMessage).where(EmailMessage.unsubscribe_token == token)
    ).scalar_one_or_none()
    if message:
        # Unsubscribing is from the customer who sent this email.
        suppress(db, message.to_email, "unsubscribed via email link", owner_id=message.owner_id)
        db.add(EmailEvent(message_id=message.id, type="unsubscribe", payload={}))
        lead = db.get(Lead, message.lead_id)
        if lead:
            lead.status = "unsubscribed"
            lead.stage = "closed"
        db.commit()
    return HTMLResponse(
        "<!doctype html><meta charset=utf-8>"
        "<title>Unsubscribed</title>"
        "<style>body{font-family:system-ui,sans-serif;max-width:32rem;margin:15vh auto;"
        "padding:0 1.5rem;color:#16202E;line-height:1.6}</style>"
        "<h1>You've been unsubscribed</h1>"
        "<p>You won't receive any further emails from us. "
        "Sorry for the interruption.</p>"
    )


@public_router.get("/api/public/bots")
def public_bots(db: Session = Depends(get_db)):
    return {
        "platform": settings.platform_name,
        "trial_hours": settings.trial_hours,
        "bots": [billing.bot_json(b) for b in billing.catalog(db)],
    }


@public_router.post("/billing/razorpay/webhook")
async def razorpay_webhook(request: Request, db: Session = Depends(get_db)):
    """Razorpay's server-to-server confirmation. Covers the customer who pays
    and closes the tab before the browser reports back."""
    import json

    body = await request.body()
    if not billing.webhook_signature_ok(body, request.headers.get("X-Razorpay-Signature", "")):
        raise HTTPException(400, "bad signature")
    event = json.loads(body or b"{}")
    kind = event.get("event", "")
    entity = ((event.get("payload") or {}).get("payment") or {}).get("entity") or {}
    order_id, payment_id = entity.get("order_id"), entity.get("id")
    if not order_id:
        return {"ok": True}
    if kind in ("payment.captured", "order.paid"):
        from app.models import Payment

        sub = billing.fulfil(db, order_id, payment_id, entity.get("amount"), entity.get("currency"))
        db.commit()
        payment = db.execute(select(Payment).where(Payment.order_id == order_id)).scalar_one_or_none()
        if sub is not None and payment is not None:
            _unlock(db, payment.user_id, payment.bot_key)
    elif kind == "payment.failed":
        from app.models import Payment

        payment = db.execute(select(Payment).where(Payment.order_id == order_id)).scalar_one_or_none()
        if payment and payment.status == "created":
            payment.status = "failed"
            payment.error = (entity.get("error_description") or "")[:500]
            db.commit()
    return {"ok": True}
